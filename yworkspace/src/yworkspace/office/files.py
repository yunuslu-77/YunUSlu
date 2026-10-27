"""Office files (Word .docx, Excel .xlsx, PowerPoint .pptx): read, create, and edit without breaking formatting.
LibreOffice (when installed) recalculates formulas, converts formats and renders pages so the AI can look at them.
These are standard Office Open XML files, so they also open in yOffice, Microsoft Office and LibreOffice."""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from ..core.errors import ToolError


def _need(module: str):
    try:
        return __import__(module)
    except ImportError as e:
        raise ToolError("no_office_libs", f"{module} is not installed.", 'pip install "yworkspace[office]"') from e


def soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def _run_soffice(args: list[str], timeout: int = 180) -> subprocess.CompletedProcess:
    exe = soffice()
    if not exe:
        raise ToolError("no_libreoffice", "LibreOffice is needed for this.", "Install LibreOffice (setup wizard).")
    profile = Path(tempfile.mkdtemp(prefix="yw-lo-"))
    try:  # a private profile so it works while the user has LibreOffice open
        return subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", *args],
                              capture_output=True, text=True, timeout=timeout)
    finally:
        shutil.rmtree(profile, ignore_errors=True)


# ============================================================================ Word
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def _add_inline(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def markdown_blocks(md: str) -> list[dict]:
    blocks, table = [], None
    for line in md.splitlines():
        s = line.rstrip()
        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            if table is None:
                table = {"type": "table", "rows": []}
                blocks.append(table)
            table["rows"].append(cells)
            continue
        table = None
        if not s.strip():
            continue
        m = re.match(r"(#{1,4})\s+(.*)", s)
        if m:
            blocks.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2)})
        elif re.match(r"\s*[-*]\s+", s):
            blocks.append({"type": "bullet", "text": re.sub(r"^\s*[-*]\s+", "", s)})
        elif re.match(r"\s*\d+[.)]\s+", s):
            blocks.append({"type": "numbered", "text": re.sub(r"^\s*\d+[.)]\s+", "", s)})
        elif s.strip() == "---pagebreak---":
            blocks.append({"type": "page_break"})
        else:
            blocks.append({"type": "paragraph", "text": s.strip()})
    return blocks


def create_docx(path: Path, blocks: list[dict] | None = None, markdown: str = "", title: str = "",
                author: str = "", base: Path | None = None) -> Path:
    docx = _need("docx")
    from docx.shared import Cm, Pt
    doc = docx.Document(str(base)) if base else docx.Document()
    if not base:
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
    doc.core_properties.title = title
    doc.core_properties.author = author or "yWorkspace"
    if title:
        doc.add_heading(title, 0)
    for b in (blocks or []) + markdown_blocks(markdown):
        kind = b.get("type", "paragraph")
        if kind == "heading":
            doc.add_heading(b["text"], min(max(int(b.get("level", 1)), 1), 4))
        elif kind == "paragraph":
            _add_inline(doc.add_paragraph(), b["text"])
        elif kind == "bullet":
            _add_inline(doc.add_paragraph(style="List Bullet"), b["text"])
        elif kind == "numbered":
            _add_inline(doc.add_paragraph(style="List Number"), b["text"])
        elif kind == "table":
            rows = b["rows"]
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    cell = t.cell(i, j)
                    cell.text = ""
                    _add_inline(cell.paragraphs[0], str(v))
                    if i == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        elif kind == "image":
            doc.add_picture(str(b["path"]), width=Cm(float(b.get("width_cm", 15))))
        elif kind == "page_break":
            doc.add_page_break()
        else:
            raise ToolError("bad_block", f"Unknown block type {kind}.",
                            "Use heading, paragraph, bullet, numbered, table, image, page_break.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _paragraphs(doc):
    """Every paragraph: body, tables (nested), headers and footers."""
    def walk_tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
                    yield from walk_tables(cell.tables)
    yield from doc.paragraphs
    yield from walk_tables(doc.tables)
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs
            yield from walk_tables(part.tables)


def read_docx(path: Path) -> str:
    docx = _need("docx")
    doc = docx.Document(str(path))
    out = []
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            style = p.style.name if p.style is not None else ""
            tag = f" ({style})" if style and style != "Normal" else ""
            out.append(f"[p{i}]{tag} {p.text}")
    for ti, t in enumerate(doc.tables):
        out.append(f"[table {ti}]")
        for r in t.rows:
            out.append("| " + " | ".join(c.text.replace("\n", " ") for c in r.cells) + " |")
    for s in doc.sections[:1]:
        h = " ".join(p.text for p in s.header.paragraphs if p.text.strip())
        f = " ".join(p.text for p in s.footer.paragraphs if p.text.strip())
        if h:
            out.append(f"[header] {h}")
        if f:
            out.append(f"[footer] {f}")
    return "\n".join(out) or "(empty document)"


def _replace_in_paragraph(p, find: str, replace: str, limit: int | None = None) -> int:
    """Replace inside a paragraph keeping the formatting of the run where each match starts."""
    count, start = 0, 0
    while limit is None or count < limit:
        text = "".join(r.text for r in p.runs)
        idx = text.find(find, start)
        if idx < 0:
            break
        end, pos, first = idx + len(find), 0, True
        for r in list(p.runs):
            a, b = pos, pos + len(r.text)
            pos = b
            if b <= idx or a >= end:
                continue
            lo, hi = max(a, idx) - a, min(b, end) - a
            if first:
                r.text = r.text[:lo] + replace + r.text[hi:]
                first = False
            else:
                r.text = r.text[:lo] + r.text[hi:]
        count += 1
        start = idx + len(replace)
    return count


def edit_docx(path: Path, edits: list[dict]) -> list[str]:
    docx = _need("docx")
    doc = docx.Document(str(path))
    paras = list(_paragraphs(doc))
    report = []
    for e in edits:
        find, rep, every = e["find"], e.get("replace", ""), bool(e.get("all"))
        hits = [p for p in paras if find in "".join(r.text for r in p.runs)]
        total = sum("".join(r.text for r in p.runs).count(find) for p in hits)
        if total == 0:
            loose = [p.text for p in paras if find.lower() in p.text.lower()]
            hint = f"Similar text exists with different capitalisation: {loose[0][:120]!r}" if loose else \
                "Read the document with office_read and copy the text exactly (one paragraph at a time)."
            raise ToolError("not_found", f"{find[:80]!r} is not in the document.", hint)
        if total > 1 and not every:
            raise ToolError("ambiguous", f"{find[:80]!r} occurs {total} times.",
                            "Include more surrounding text, or set all=true to change every occurrence.")
        done = 0
        for p in hits:
            done += _replace_in_paragraph(p, find, rep, None if every else 1)
            if done and not every:
                break
        report.append(f"{find[:40]!r} → {rep[:40]!r}: {done}×")
    doc.save(str(path))
    return report


# ============================================================================ Excel
def read_xlsx(path: Path, sheet: str = "", cell_range: str = "", limit: int = 200) -> str:
    openpyxl = _need("openpyxl")
    wb_f = openpyxl.load_workbook(str(path), data_only=False)
    wb_v = openpyxl.load_workbook(str(path), data_only=True)
    names = [sheet] if sheet else wb_f.sheetnames
    out = [f"Sheets: {', '.join(wb_f.sheetnames)}"]
    for name in names:
        if name not in wb_f.sheetnames:
            raise ToolError("no_sheet", f"No sheet {name!r}.", f"Sheets: {', '.join(wb_f.sheetnames)}")
        wf, wv = wb_f[name], wb_v[name]
        out.append(f"## {name} ({wf.dimensions})")
        cells = wf[cell_range] if cell_range else wf.iter_rows()
        if cell_range and not isinstance(cells, tuple):
            cells = ((cells,),)
        elif cell_range and cells and not isinstance(cells[0], tuple):
            cells = (cells,)
        n = 0
        for row in cells:
            for c in row:
                if c.value is None:
                    continue
                v = wv[c.coordinate].value
                if isinstance(c.value, str) and c.value.startswith("="):
                    out.append(f"{c.coordinate}: {c.value}  → {v if v is not None else '(not calculated)'}")
                else:
                    out.append(f"{c.coordinate}: {c.value!r}")
                n += 1
                if n >= limit:
                    out.append("… more cells (narrow with cell_range)")
                    return "\n".join(out)
    return "\n".join(out)


def recalculate(path: Path) -> bool:
    """Open and save through LibreOffice so formula results are stored in the file."""
    if not soffice():
        return False
    out = Path(tempfile.mkdtemp(prefix="yw-calc-"))
    try:
        r = _run_soffice(["--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(path)])
        produced = out / (path.stem + ".xlsx")
        if r.returncode != 0 or not produced.exists():
            return False
        shutil.copyfile(produced, path)
        return True
    finally:
        shutil.rmtree(out, ignore_errors=True)


ERRORS = ("#DIV/0!", "#NAME?", "#REF!", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


def write_xlsx(path: Path, sheet: str, cells: dict, widths: dict | None = None, bold_rows: list[int] | None = None,
               number_formats: dict | None = None) -> tuple[list[str], bool]:
    """cells: {"A1": "Ürün", "B2": 12.5, "B10": "=SUM(B2:B9)"}. Returns (formula results, recalculated)."""
    openpyxl = _need("openpyxl")
    from openpyxl.styles import Font
    path = Path(path)
    if path.exists():
        wb = openpyxl.load_workbook(str(path))
    else:
        wb = openpyxl.Workbook()
        wb.active.title = sheet or "Sheet1"
    ws = wb[sheet] if sheet in wb.sheetnames else (wb.create_sheet(sheet) if sheet else wb.active)
    for ref, value in cells.items():
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", ref):
            raise ToolError("bad_cell", f"{ref} is not a cell reference like A1.")
        ws[ref] = value
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = float(w)
    for r in bold_rows or []:
        for c in ws[int(r)]:
            c.font = Font(bold=True)
    for ref, fmt in (number_formats or {}).items():
        ws[ref].number_format = fmt
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    formulas = [r for r, v in cells.items() if isinstance(v, str) and v.startswith("=")]
    recalculated = recalculate(path) if formulas else False
    results = []
    if formulas:
        wv = openpyxl.load_workbook(str(path), data_only=True)[ws.title]
        for r in formulas:
            results.append(f"{r}: {cells[r]} → {wv[r].value!r}")
    return results, recalculated


# ============================================================================ PowerPoint
def create_pptx(path: Path, slides: list[dict], title: str = "", template: Path | None = None) -> list[str]:
    pptx = _need("pptx")
    from pptx.util import Inches, Pt
    prs = pptx.Presentation(str(template)) if template else pptx.Presentation()
    if not template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = title
    warnings = []
    for i, s in enumerate(slides, 1):
        layout = s.get("layout", "content")
        idx = {"title": 0, "content": 1, "section": 2, "title_only": 5, "blank": 6, "image": 5}.get(layout, 1)
        slide = prs.slides.add_slide(prs.slide_layouts[idx])
        if not template:  # the built-in layouts are 4:3; stretch their boxes to the 16:9 page
            for ph in slide.placeholders:
                top, height = ph.top, ph.height  # inherited from the layout; must be set together
                ph.left, ph.top, ph.width, ph.height = Inches(0.7), top, prs.slide_width - Inches(1.4), height
        if slide.shapes.title is not None:
            slide.shapes.title.text = s.get("title", "")
        bullets = s.get("bullets", [])
        if layout == "title" and len(slide.placeholders) > 1:
            slide.placeholders[1].text = s.get("subtitle", "")
        elif bullets and len(slide.placeholders) > 1:
            body = slide.placeholders[1].text_frame
            body.clear()
            for j, b in enumerate(bullets):
                level = 1 if b.startswith("  ") else 0
                para = body.paragraphs[0] if j == 0 else body.add_paragraph()
                para.text = b.strip()
                para.level = level
                for r in para.runs:
                    r.font.size = Pt(24 if level == 0 else 20)
        if s.get("image"):
            top = Inches(1.6)
            slide.shapes.add_picture(str(s["image"]), Inches(1), top, height=prs.slide_height - top - Inches(0.5))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        if len(s.get("title", "")) > 60:
            warnings.append(f"slide {i}: title is {len(s['title'])} characters - say the message in fewer words")
        if len(bullets) > 6:
            warnings.append(f"slide {i}: {len(bullets)} bullets - split the slide (max 6)")
        long = [b for b in bullets if len(b) > 90]
        if long:
            warnings.append(f"slide {i}: {len(long)} bullets longer than 90 characters - move detail to the notes")
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return warnings


def read_pptx(path: Path) -> str:
    pptx = _need("pptx")
    prs = pptx.Presentation(str(path))
    out = [f"{len(prs.slides)} slides"]
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip().replace("\n", " / "))
            elif sh.shape_type == 13:
                texts.append("[picture]")
            elif getattr(sh, "has_table", False) and sh.has_table:
                texts.append("[table]")
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        out.append(f"--- slide {i}: " + " | ".join(texts) + (f"\n    notes: {notes}" if notes else ""))
    return "\n".join(out)


# ============================================================================ convert and render
def convert(path: Path, fmt: str, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    src_ext = path.suffix.lower()
    if src_ext == ".md" and fmt in ("docx", "odt", "html") and shutil.which("pandoc"):
        dest = out_dir / f"{path.stem}.{fmt}"
        r = subprocess.run(["pandoc", str(path), "-o", str(dest)], capture_output=True, text=True, timeout=120)
        if r.returncode != 0:
            raise ToolError("convert_failed", r.stderr[-1000:])
        return dest
    if fmt == "md" and src_ext in (".docx", ".odt", ".html") and shutil.which("pandoc"):
        dest = out_dir / f"{path.stem}.md"
        r = subprocess.run(["pandoc", str(path), "-t", "gfm", "-o", str(dest)], capture_output=True, text=True,
                           timeout=120)
        if r.returncode != 0:
            raise ToolError("convert_failed", r.stderr[-1000:])
        return dest
    r = _run_soffice(["--convert-to", fmt, "--outdir", str(out_dir), str(path)])
    dest = out_dir / f"{path.stem}.{fmt.split(':')[0]}"
    if r.returncode != 0 or not dest.exists():
        raise ToolError("convert_failed", (r.stdout + r.stderr)[-1000:] or "LibreOffice produced no file.")
    return dest


def render_pages(path: Path, out_dir: Path, max_pages: int = 20, dpi: int = 60) -> list[Path]:
    """Any office file → PNG per page/slide, so the AI (or a vision model) can check how it looks."""
    pdf = path if path.suffix.lower() == ".pdf" else convert(path, "pdf", out_dir)
    if not shutil.which("pdftoppm"):
        raise ToolError("no_pdftoppm", "pdftoppm (poppler-utils) is needed to render pages.")
    prefix = out_dir / path.stem
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-l", str(max_pages), str(pdf), str(prefix)],
                   check=True, capture_output=True, timeout=300)
    return sorted(out_dir.glob(f"{path.stem}-*.png"))
