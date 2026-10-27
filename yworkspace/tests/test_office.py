import shutil

import pytest

pytest.importorskip("docx")
pytest.importorskip("pptx")
pytest.importorskip("openpyxl")

from yworkspace.app import build_session  # noqa: E402
from yworkspace.config import Config  # noqa: E402

needs_lo = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")


def session(tmp_path, monkeypatch):
    monkeypatch.setenv("YW_HOME", str(tmp_path / "home"))
    return build_session(tmp_path / "ofis", Config(sandbox="local"), with_engine=False)


def test_word_create_read_and_precise_edit(tmp_path, monkeypatch):
    import docx
    sess = session(tmp_path, monkeypatch)
    call = sess.registry.call
    root = sess.workspace.root
    r = call("office_docx_create", {"path": "teklif.docx", "title": "Fiyat Teklifi", "markdown":
                                    "## Kapsam\nRandevu sistemi için **toplam bedel 45.000 TL** olarak belirlenmiştir.\n"
                                    "- Web paneli\n- Android uygulaması\n\n| Kalem | Tutar |\n|---|---|\n"
                                    "| Web | 30.000 TL |\n| Android | 15.000 TL |\n"})
    assert not r.is_error, r.text
    text = call("office_read", {"path": "teklif.docx"}).text
    assert "(Heading 2) Kapsam" in text and "| Android | 15.000 TL |" in text and "(List Bullet) Web paneli" in text
    # A run-split phrase: "toplam bedel 45.000 TL" is bold inside a normal paragraph.
    e = call("office_docx_edit", {"path": "teklif.docx", "edits": [
        {"find": "toplam bedel 45.000 TL", "replace": "toplam bedel 48.000 TL"},
        {"find": "15.000 TL", "replace": "18.000 TL"}]})
    assert not e.is_error and "previous version saved" in e.text, e.text
    d = docx.Document(str(root / "teklif.docx"))
    par = next(p for p in d.paragraphs if "48.000" in p.text)
    assert par.text == "Randevu sistemi için toplam bedel 48.000 TL olarak belirlenmiştir."
    assert any(r.bold and "48.000" in r.text for r in par.runs) and not par.runs[0].bold
    assert d.tables[0].cell(2, 1).text == "18.000 TL"
    amb = call("office_docx_edit", {"path": "teklif.docx", "edits": [{"find": "TL", "replace": "₺"}]})
    assert amb.is_error and "ambiguous" in amb.text
    every = call("office_docx_edit", {"path": "teklif.docx", "edits": [{"find": "TL", "replace": "TL (KDV hariç)",
                                                                         "all": True}]}).text
    assert "3×" in every
    miss = call("office_docx_edit", {"path": "teklif.docx", "edits": [{"find": "TOPLAM BEDEL", "replace": "x"}]})
    assert miss.is_error and "capitalisation" in miss.text
    assert len(list((root / ".yw/office-backup").glob("teklif.*.docx"))) >= 2


@needs_lo
def test_excel_formulas_are_calculated(tmp_path, monkeypatch):
    sess = session(tmp_path, monkeypatch)
    call = sess.registry.call
    r = call("office_xlsx_write", {"path": "butce.xlsx", "sheet": "Ekim", "bold_rows": [1], "widths": {"A": 20},
                                   "cells": {"A1": "Kalem", "B1": "Tutar", "A2": "Kira", "B2": 15000, "A3": "Market",
                                             "B3": 6250.5, "A4": "Toplam", "B4": "=SUM(B2:B3)", "C4": "=B4/0"}})
    assert "B4: =SUM(B2:B3) → 21250.5" in r.text, r.text
    assert "FORMULA ERRORS" in r.text and "C4" in r.text.split("FORMULA ERRORS")[1]
    read = call("office_read", {"path": "butce.xlsx", "sheet": "Ekim"}).text
    assert "B4: =SUM(B2:B3)  → 21250.5" in read and "A2: 'Kira'" in read


@needs_lo
def test_powerpoint_render_and_convert(tmp_path, monkeypatch):
    sess = session(tmp_path, monkeypatch)
    call = sess.registry.call
    root = sess.workspace.root
    r = call("office_pptx_create", {"path": "sunum.pptx", "title": "Randevum", "slides": [
        {"layout": "title", "title": "Randevum", "subtitle": "Kuaförler için randevu"},
        {"title": "Neden?", "bullets": ["Telefon trafiği azalır", "  Gece de randevu alınır", "Hatırlatma SMS'i"],
         "notes": "Rakamları anlat"},
        {"title": "Çok uzun bir başlık " * 5, "bullets": [f"madde {i}" for i in range(8)]}]})
    assert "slide 3: title is" in r.text and "8 bullets" in r.text
    read = call("office_read", {"path": "sunum.pptx"}).text
    assert "3 slides" in read and "notes: Rakamları anlat" in read and "Gece de randevu alınır" in read
    shot = call("office_render", {"path": "sunum.pptx"})
    assert not shot.is_error and len(shot.images) == 3, shot.text
    pdf = call("office_convert", {"path": "sunum.pptx", "to": "pdf"}).text
    assert "sunum.pdf" in pdf and (root / "sunum.pdf").stat().st_size > 1000
    (root / "not.md").write_text("# Başlık\n\nMerhaba **dünya**\n")
    call("office_convert", {"path": "not.md", "to": "docx"})
    assert "(Title) Başlık" in call("office_read", {"path": "not.docx"}).text or \
        "Başlık" in call("office_read", {"path": "not.docx"}).text
